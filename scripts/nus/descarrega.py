#!/usr/bin/env python3
"""Descarrega les fonts de «El nus» i les desa a data/raw/nus-*.csv.

Ús:  python3 scripts/nus/descarrega.py              (descarrega dels servidors oficials; cal `pip install xlrd openpyxl`)
     python3 scripts/nus/descarrega.py --local DIR  (llegeix els mateixos fitxers originals des d'una carpeta)
Després: node scripts/nus/dades.mjs

Amb --local, DIR ha de contenir els fitxers amb el nom que indica FILES (els originals, sense modificar).
"""
import csv, io, os, re, sys, urllib.request, collections

RAW = os.path.join(os.path.dirname(__file__), '..', '..', 'data', 'raw')
UA = {'User-Agent': 'Mozilla/5.0 (terra-i-gent)'}
ESTAT = 'https://ec.europa.eu/eurostat/api/dissemination/sdmx/2.1/data/ilc_lvho07c/A.PC.{t}.ES+EU27_2020?format=SDMX-CSV'
TEN = {'OWN_ML': 'propietari amb hipoteca', 'OWN_NL': 'propietari sense hipoteca', 'RENT_MKT': 'lloguer de mercat', 'RENT_FR': 'lloguer reduït o gratuït'}
FILES = {
    # Ministeri d'Habitatge i Agenda Urbana, Boletín estadístico: taules 3.1, 3.2 i 1.6
    'mitms_32200500.xls': 'https://apps.fomento.gob.es/BoletinOnline2/sedal/32200500.XLS',
    'mitms_32201000.xls': 'https://apps.fomento.gob.es/BoletinOnline2/sedal/32201000.XLS',
    'mitms_31306000.xls': 'https://apps.fomento.gob.es/BoletinOnline2/sedal/31306000.XLS',
    # Banc d'Espanya, Síntesi d'indicadors, quadre 1.5 (habitatge)
    'bde_si_1_5.csv': 'https://www.bde.es/webbe/es/estadisticas/compartido/datos/csv/si_1_5.csv',
    # OCDE, Affordable Housing Database
    'PH4-2-Social-rental-housing-stock.xlsx': 'https://webfs.oecd.org/Els-com/Affordable_Housing_Database/PH4-2-Social-rental-housing-stock.xlsx',
    'HM1-3-Housing-tenures.xlsx': 'https://webfs.oecd.org/Els-com/Affordable_Housing_Database/HM1-3-Housing-tenures.xlsx',
    # BIS, Selected residential property prices (Espanya, nominal i real, 2010 = 100)
    'bis_spp.csv': 'https://stats.bis.org/api/v1/data/BIS,WS_SPP,1.0/Q.ES..628?format=csv',
    # Eurostat ilc_lvho07c: amb --local, l'exportació en etiquetes (estat_ilc_lvho07c.csv); en línia, l'API SDMX-CSV
    'estat_ilc_lvho07c.csv': None,
}
LOCAL = sys.argv[sys.argv.index('--local') + 1] if '--local' in sys.argv else None

def get(url):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=180) as r:
        return r.read()

def src(name):
    if LOCAL:
        with open(os.path.join(LOCAL, name), 'rb') as f: return f.read()
    return get(FILES[name])

def save(name, header, rows):
    with open(os.path.join(RAW, f'nus-{name}.csv'), 'w', newline='') as f:
        w = csv.writer(f, lineterminator='\n'); w.writerow(header); w.writerows(rows)
    print(f'nus-{name}.csv', len(rows), 'files')

# --- Habitatges iniciats i acabats (Ministeri). Taula 3.1 lliures iniciats, 3.2 lliures acabats, 1.6 protegits amb qualificació definitiva.
import xlrd
def mitms(name):
    s = xlrd.open_workbook(file_contents=src(name)).sheet_by_index(0)
    hdr = next(r for r in range(s.nrows) if str(s.cell_value(r, 2)).startswith('19'))
    years = [int(float(v)) for v in s.row_values(hdr)[2:] if str(v).strip()]
    out = {}
    for r in range(hdr + 1, s.nrows):
        k = str(s.cell_value(r, 1)).strip()
        if k in ('TOTAL NACIONAL', 'Cataluña', 'Barcelona'):
            out[k] = {y: (int(float(v)) if str(v).strip() not in ('', '-') else None) for y, v in zip(years, s.row_values(r)[2:])}
    return out
ini, aca, vpo = mitms('mitms_32200500.xls'), mitms('mitms_32201000.xls'), mitms('mitms_31306000.xls')
ys = sorted(aca['Cataluña'])
save('habitatges', ['any', 'cat_iniciats_lliures', 'cat_acabats_lliures', 'cat_protegits', 'bcn_iniciats_lliures', 'bcn_acabats_lliures', 'bcn_protegits', 'es_iniciats_lliures', 'es_acabats_lliures', 'es_protegits'],
     [[y, ini['Cataluña'][y], aca['Cataluña'][y], vpo['Cataluña'][y], ini['Barcelona'][y], aca['Barcelona'][y], vpo['Barcelona'][y],
       ini['TOTAL NACIONAL'][y], aca['TOTAL NACIONAL'][y], vpo['TOTAL NACIONAL'][y]] for y in ys])

# --- Banc d'Espanya, quadre 1.5: mitjanes anuals de les sèries mensuals o trimestrals (només anys complets).
MES = {'ENE': 1, 'FEB': 2, 'MAR': 3, 'ABR': 4, 'MAY': 5, 'JUN': 6, 'JUL': 7, 'AGO': 8, 'SEP': 9, 'OCT': 10, 'NOV': 11, 'DIC': 12}
raw = src('bde_si_1_5.csv')
try: txt = raw.decode('utf-8')
except UnicodeDecodeError: txt = raw.decode('latin-1')
B = list(csv.reader(io.StringIO(txt.lstrip('﻿'))))
row = next(r for r in B if r and r[0].startswith('ALIAS'))
col = {a: j for j, a in enumerate(row)}
frow = next(r for r in B if r and r[0].startswith('FRECUENCIA'))
freq = {a: frow[j] for a, j in col.items()}
def annual(a, how='mean'):
    j = col[a]; n = {'MENSUAL': 12, 'DIARIA': 12, 'TRIMESTRAL': 4, 'ANUAL': 1, 'SEMESTRAL': 2}.get(freq[a], 12); g = collections.defaultdict(list)
    for r in B:
        m = re.match(r'([A-Z]{3}) (\d{4})', r[0] if r else '')
        if not m or j >= len(r): continue
        v = r[j].strip()
        if v in ('', '_', '-'): continue
        g[int(m.group(2))].append(float(v.replace(',', '.')))
    return {y: (round(sum(v) / len(v), 3) if how == 'mean' else round(sum(v))) for y, v in g.items() if len(v) == n}
BDE = [('llars_milers', 'SI_1_5.36', 'mean'), ('tipus_hipoteca', 'SI_1_5.40', 'mean'), ('preu_renda_anys', 'SI_1_5.43', 'mean'),
       ('esforc_teoric', 'SI_1_5.44', 'mean'), ('rendibilitat_lloguer', 'SI_1_5.63', 'mean'), ('diposit_1_2_anys', 'SI_1_5.65', 'mean'),
       ('credit_habitatge_pib', 'SI_1_5.56', 'mean'), ('cost_construccio', 'SI_1_5.14', 'mean'),
       ('pct_lloguer', 'SI_1_5.91', 'mean'), ('pct_propietat', 'SI_1_5.90', 'mean')]
series = {k: annual(a, h) for k, a, h in BDE}
ys = sorted(set().union(*[s.keys() for s in series.values()]))
save('bde', ['any'] + [k for k, _, _ in BDE], [[y] + [series[k].get(y) for k, _, _ in BDE] for y in ys])

# --- OCDE: parc de lloguer social (PH4.2, taula A1) i tinença per edat (HM1.3, taula A6, Espanya).
import openpyxl
def wb(name): return openpyxl.load_workbook(io.BytesIO(src(name)), read_only=True, data_only=True)
ws = wb('PH4-2-Social-rental-housing-stock.xlsx')['Table PH4.2.A1']
rows = []
for r in ws.iter_rows(min_row=6, values_only=True):
    if not r[0] or str(r[0]).startswith('Notes'): break
    c = list(r) + [None] * 8
    name = re.sub(r'\s*\(\d+\)\s*$', '', str(c[0])).strip()
    # Columnes: [país, n 2010, % 2010, any 2010, n 2022, % 2022, any 2022]
    if c[5] not in (None, '') and c[6] not in (None, ''):
        rows.append([name, int(float(c[6])), round(float(c[5]), 2)])
save('social', ['pais', 'any', 'pct_parc'], rows)
ws = wb('HM1-3-Housing-tenures.xlsx')['HM1.3.A6']
grid = [list(r) for r in ws.iter_rows(values_only=True)]
ages, years = grid[3], grid[4]
blocks, cur = [], None
for j, v in enumerate(ages):
    if v: cur = str(v).strip()
    if cur and j >= 2 and years[j] not in (None, ''): blocks.append((j, cur, int(float(years[j]))))
rows, country = [], None
for r in grid[5:]:
    if r[0]: country = str(r[0]).strip()
    if country != 'Spain' or not r[1]: continue
    for j, age, y in blocks:
        v = r[j]
        if v not in (None, '..', ''): rows.append([age, str(r[1]).strip(), y, round(float(v), 2)])
save('tinenca', ['edat', 'tinenca', 'any', 'pct_persones'], rows)

# --- BIS: preus residencials d'Espanya, nominal (N) i real (R), mitjana anual de l'índex trimestral 2010 = 100.
g = collections.defaultdict(lambda: collections.defaultdict(list))
for r in csv.DictReader(io.StringIO(src('bis_spp.csv').decode('utf-8-sig'))):
    if r['REF_AREA'] == 'ES' and r['UNIT_MEASURE'] == '628' and r['OBS_VALUE']:
        g[r['VALUE']][int(r['TIME_PERIOD'][:4])].append(float(r['OBS_VALUE']))
ys = sorted(y for y in g['R'] if len(g['R'][y]) == 4)
save('bis', ['any', 'nominal_2010_100', 'real_2010_100'], [[y, round(sum(g['N'][y]) / 4, 2), round(sum(g['R'][y]) / 4, 2)] for y in ys])

# --- Eurostat ilc_lvho07c: població en llars amb despeses d'habitatge superiors al 40% de la renda disponible, per tinença.
rows = []
if LOCAL:
    L = {'Owner, with mortgage or loan': 'OWN_ML', 'Owner, no outstanding mortgage or housing loan': 'OWN_NL',
         'Tenant, rent at market price': 'RENT_MKT', 'Tenant, rent at reduced price or free': 'RENT_FR'}
    G = {'Spain': 'ES', 'European Union - 27 countries (from 2020)': 'EU27_2020'}
    rd = list(csv.reader(io.StringIO(src('estat_ilc_lvho07c.csv').decode('utf-8-sig'))))
    yrs = [int(float(v)) for v in rd[0][4:]]
    for r in rd[1:]:
        if r[2] in L and r[3] in G:
            for y, v in zip(yrs, r[4:]):
                v = v.split(' ')[0]
                if v not in (':', ''): rows.append([G[r[3]], L[r[2]], y, float(v)])
else:
    for t in TEN:
        for r in csv.DictReader(io.StringIO(get(ESTAT.format(t=t)).decode('utf-8'))):
            if r['OBS_VALUE']: rows.append([r['geo'], t, int(r['TIME_PERIOD']), float(r['OBS_VALUE'])])
save('sobrecarrega', ['geo', 'tinenca', 'any', 'pct'], sorted(rows))
